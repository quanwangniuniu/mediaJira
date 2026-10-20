import io
import tempfile
from contextlib import contextmanager
from unittest.mock import patch

from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from rest_framework.test import APIClient

from spreadsheet.import_baseline import build_import_baseline
from spreadsheet.models import (
    Cell,
    CellValueType,
    SheetColumn,
    SheetImportSource,
    SheetRow,
    SpreadsheetCellFormat,
    SpreadsheetHighlight,
)
from spreadsheet.services import SheetService
from spreadsheet.tests.test_services import (
    create_test_organization,
    create_test_project,
    create_test_sheet,
    create_test_spreadsheet,
    create_test_user,
)
from spreadsheet.xlsx_export import build_sheet_workbook
from spreadsheet.xlsx_patch_export import XlsxPatchExportError


def _workbook_bytes(build):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Sheet1'
    build(workbook, worksheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@contextmanager
def _open_exported(content):
    workbook = load_workbook(io.BytesIO(content))
    try:
        yield workbook
    finally:
        workbook.close()


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class TestXlsxPatchExport(TestCase):
    def setUp(self):
        self.user = create_test_user()
        self.organization = create_test_organization()
        self.project = create_test_project(self.organization, owner=self.user)
        self.spreadsheet = create_test_spreadsheet(self.project)
        self.sheet = create_test_sheet(self.spreadsheet)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.url = (
            f'/api/spreadsheet/spreadsheets/{self.spreadsheet.slug}/'
            f'sheets/{self.sheet.id}/export.xlsx'
        )

    def _cell(self, row_pos, col_pos, raw):
        row, _created = SheetRow.objects.get_or_create(sheet=self.sheet, position=row_pos)
        column, _created = SheetColumn.objects.get_or_create(
            sheet=self.sheet,
            position=col_pos,
            defaults={'name': chr(ord('A') + col_pos)},
        )
        return Cell.objects.create(
            sheet=self.sheet,
            row=row,
            column=column,
            raw_input=raw,
            string_value=raw,
            value_type=CellValueType.STRING,
        )

    def _save_excel_source(self, original_bytes, worksheet_name='Sheet1'):
        """Call only after the sheet matches the workbook and before any edit."""
        return SheetImportSource.objects.create(
            sheet=self.sheet,
            origin=SheetImportSource.Origin.EXCEL,
            worksheet_name=worksheet_name,
            original_filename='book.xlsx',
            baseline=build_import_baseline(self.sheet),
            file=ContentFile(original_bytes, name='book.xlsx'),
        )

    def _assert_rebuild(self):
        captured = {}
        real = build_sheet_workbook

        def wrapper(sheet):
            captured['content'] = real(sheet)
            return captured['content']

        with patch('spreadsheet.views.build_sheet_workbook', side_effect=wrapper):
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, captured['content'])

    def test_no_source_uses_rebuild(self):
        self._cell(0, 0, 'plain')
        self._assert_rebuild()

    def test_google_source_uses_rebuild(self):
        self._cell(0, 0, 'plain')
        SheetImportSource.objects.create(
            sheet=self.sheet,
            origin=SheetImportSource.Origin.GOOGLE,
            worksheet_name='Budget',
            google_spreadsheet_id='abc',
            google_sheet_id=0,
        )
        self._assert_rebuild()

    def test_no_changes_return_original_bytes(self):
        self._cell(0, 0, 'same')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'same'))
        self._save_excel_source(original)
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, original)

    def test_value_change_patches_only_that_cell(self):
        changed = self._cell(0, 0, 'old')
        self._cell(0, 1, 'keep')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'old'),
                worksheet.__setitem__('B1', 'keep'),
            )
        )
        self._save_excel_source(original)
        changed.raw_input = 'new'
        changed.string_value = 'new'
        changed.save(update_fields=['raw_input', 'string_value'])
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            worksheet = workbook['Sheet1']
            self.assertEqual(worksheet['A1'].value, 'new')
            self.assertEqual(worksheet['B1'].value, 'keep')

    def test_bold_change_keeps_underline(self):
        self._cell(0, 0, 'styled')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'styled'),
                setattr(worksheet['A1'], 'font', Font(bold=False, underline='single')),
            )
        )
        self._save_excel_source(original)
        SpreadsheetCellFormat.objects.create(
            sheet=self.sheet, row_index=0, column_index=0, bold=True,
        )
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.font.underline, 'single')

    def test_highlight_change_keeps_border_and_alignment(self):
        self._cell(0, 0, 'boxed')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'boxed'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                setattr(worksheet['A1'], 'border', Border(left=Side(style='thin'))),
                setattr(worksheet['A1'], 'alignment', Alignment(horizontal='center')),
            )
        )
        self._save_excel_source(original)
        SpreadsheetHighlight.objects.create(
            spreadsheet=self.spreadsheet,
            sheet=self.sheet,
            scope='CELL',
            row_index=0,
            col_index=0,
            color='#BFDBFE',
        )
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertIn('BFDBFE', cell.fill.fgColor.rgb)
            self.assertEqual(cell.font.underline, 'single')
            self.assertEqual(cell.border.left.style, 'thin')
            self.assertEqual(cell.alignment.horizontal, 'center')

    def test_unchanged_custom_number_format_survives_bold(self):
        self._cell(0, 0, '10')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 10),
                setattr(worksheet['A1'], 'number_format', '[$AUD]#,##0.00'),
            )
        )
        self._save_excel_source(original)
        SpreadsheetCellFormat.objects.create(
            sheet=self.sheet, row_index=0, column_index=0, bold=True,
        )
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            cell = workbook['Sheet1']['A1']
            self.assertTrue(cell.font.bold)
            self.assertEqual(cell.number_format, '[$AUD]#,##0.00')

    def test_other_worksheet_is_preserved(self):
        self._cell(0, 0, 'old')

        def build(workbook, worksheet):
            worksheet['A1'] = 'old'
            other = workbook.create_sheet('Sheet2')
            other['A1'] = 'other-sheet'

        original = _workbook_bytes(build)
        source_cell = Cell.objects.get(sheet=self.sheet)
        self._save_excel_source(original)
        source_cell.raw_input = 'new'
        source_cell.string_value = 'new'
        source_cell.save(update_fields=['raw_input', 'string_value'])
        response = self.client.get(self.url)
        with _open_exported(response.content) as workbook:
            self.assertEqual(workbook.sheetnames, ['Sheet1', 'Sheet2'])
            self.assertEqual(workbook['Sheet1']['A1'].value, 'new')
            self.assertEqual(workbook['Sheet2']['A1'].value, 'other-sheet')

    def test_column_insert_uses_rebuild(self):
        self._cell(0, 0, 'left')
        self._cell(0, 1, 'right')
        original = _workbook_bytes(
            lambda workbook, worksheet: (
                worksheet.__setitem__('A1', 'left'),
                setattr(worksheet['A1'], 'font', Font(underline='single')),
                worksheet.__setitem__('B1', 'right'),
            )
        )
        self._save_excel_source(original)
        SheetService.insert_columns(self.sheet, position=1, count=1)
        self._assert_rebuild()

    def test_patch_failure_falls_back_to_rebuild(self):
        cell = self._cell(0, 0, 'old')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'old'))
        self._save_excel_source(original)
        cell.raw_input = 'new'
        cell.string_value = 'new'
        cell.save(update_fields=['raw_input', 'string_value'])
        with patch(
            'spreadsheet.views.build_patched_workbook',
            side_effect=XlsxPatchExportError('patch failed'),
        ):
            self._assert_rebuild()

    def test_source_read_failure_falls_back_to_rebuild(self):
        self._cell(0, 0, 'same')
        original = _workbook_bytes(lambda workbook, worksheet: worksheet.__setitem__('A1', 'same'))
        self._save_excel_source(original)
        with patch(
            'django.db.models.fields.files.FieldFile.open',
            side_effect=OSError('unreadable'),
        ):
            self._assert_rebuild()