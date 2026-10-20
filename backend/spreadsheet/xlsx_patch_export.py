"""Patch a copy of an imported .xlsx from a Phase 3 diff.

This module does not rebuild a workbook and does not catch its own failures.
The export view decides whether to fall back to build_sheet_workbook.
"""

import io
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.styles.colors import Color

from .models import Cell, SheetColumn, SheetRow
from .sparkline import is_sparkline
from .xlsx_export import (
    _EXCEL_FONT_NAMES,
    _apply_fill,
    _cell_export_value,
    _excel_number_format,
    _excel_rgb,
)

# _excel_font omits name, size, and color when the web value is null.
# A brand-new openpyxl cell then uses these library defaults. This repo does
# not define its own Excel font name, point size, or cleared text color.
# The grid draws a null font size at 12px; the Excel exporter does not write 12.
_OPENPYXL_DEFAULT_FONT_NAME = 'Calibri'
_OPENPYXL_DEFAULT_FONT_SIZE = 11
_OPENPYXL_DEFAULT_FONT_COLOR = Color(theme=1)

_FONT_FIELDS = (
    'bold',
    'italic',
    'strikethrough',
    'text_color',
    'font_family',
    'font_size',
)


class XlsxPatchExportError(Exception):
    """The original workbook cannot be patched safely."""


def build_patched_workbook(source, diff) -> bytes:
    """Return xlsx bytes for source.file with diff applied to its worksheet."""
    if not getattr(source, 'file', None):
        raise XlsxPatchExportError('The original Excel file is missing.')

    sheet = source.sheet
    try:
        with source.file.open('rb') as handle:
            workbook = load_workbook(handle)
    except XlsxPatchExportError:
        raise
    except Exception as exc:
        raise XlsxPatchExportError('The original Excel file could not be opened.') from exc

    try:
        worksheet = _worksheet(workbook, getattr(source, 'worksheet_name', ''))
        _reject_sparkline_changes(sheet, diff.get('cells') or {})
        for key, change in (diff.get('cells') or {}).items():
            _patch_value(worksheet, sheet, key, change)
        for key, change in (diff.get('formats') or {}).items():
            _patch_format(worksheet, key, change or {})
        highlights = diff.get('highlights') or {}
        for key, color in (highlights.get('cells') or {}).items():
            _patch_highlight(worksheet, key, color)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()


def _worksheet(workbook, worksheet_name):
    name = (worksheet_name or '').strip()
    if not name or name not in workbook.sheetnames:
        raise XlsxPatchExportError(
            f'Worksheet {worksheet_name!r} is not in the original workbook.'
        )
    return workbook[name]


def _parse_coordinate(key):
    parts = str(key).split(',')
    if len(parts) != 2:
        raise XlsxPatchExportError(f'Bad cell coordinate {key!r}.')
    try:
        row_pos = int(parts[0])
        col_pos = int(parts[1])
    except ValueError as exc:
        raise XlsxPatchExportError(f'Bad cell coordinate {key!r}.') from exc
    if row_pos < 0 or col_pos < 0:
        raise XlsxPatchExportError(f'Bad cell coordinate {key!r}.')
    return row_pos, col_pos


def _active_at(model, sheet, position, label):
    obj = model.objects.filter(sheet=sheet, position=position, is_deleted=False).first()
    if obj is None:
        raise XlsxPatchExportError(f'No active {label} at position {position}.')
    return obj


def _live_cell(sheet, row, column):
    return (
        Cell.objects.filter(sheet=sheet, row=row, column=column, is_deleted=False)
        .first()
    )


def _excel_cell(worksheet, key):
    row_pos, col_pos = _parse_coordinate(key)
    return worksheet.cell(row=row_pos + 1, column=col_pos + 1)


def _reject_sparkline_changes(sheet, cell_changes):
    for key, change in cell_changes.items():
        raw = (change or {}).get('raw_input')
        if raw is None:
            continue
        if is_sparkline(str(raw)):
            raise XlsxPatchExportError(f'Cell {key} is a sparkline and cannot be patched.')
        row_pos, col_pos = _parse_coordinate(key)
        row = _active_at(SheetRow, sheet, row_pos, 'row')
        column = _active_at(SheetColumn, sheet, col_pos, 'column')
        cell = _live_cell(sheet, row, column)
        if cell is not None and is_sparkline(cell.raw_input):
            raise XlsxPatchExportError(f'Cell {key} is a sparkline and cannot be patched.')


def _patch_value(worksheet, sheet, key, change):
    row_pos, col_pos = _parse_coordinate(key)
    excel_cell = worksheet.cell(row=row_pos + 1, column=col_pos + 1)
    raw = (change or {}).get('raw_input')
    if raw is None:
        excel_cell.value = None
        return

    row = _active_at(SheetRow, sheet, row_pos, 'row')
    column = _active_at(SheetColumn, sheet, col_pos, 'column')
    cell = _live_cell(sheet, row, column)
    if cell is None:
        raise XlsxPatchExportError(f'No active cell at {key} to convert for export.')

    stored = (cell.raw_input or '').strip()
    if stored.startswith('='):
        excel_cell.value = stored
        return
    excel_cell.value = _cell_export_value(cell)


def _apply_font_family(font, value) -> bool:
    """Return True when the font name was changed."""
    if value is None:
        font.name = _OPENPYXL_DEFAULT_FONT_NAME
        font.scheme = None
        return True
    mapped = _EXCEL_FONT_NAMES.get(str(value).strip())
    if not mapped:
        return False
    font.name = mapped
    font.scheme = None
    return True


def _patch_font(excel_cell, changes):
    if not any(field in changes for field in _FONT_FIELDS):
        return
    font = copy(excel_cell.font)
    changed = False
    if 'bold' in changes:
        font.bold = bool(changes['bold'])
        changed = True
    if 'italic' in changes:
        font.italic = bool(changes['italic'])
        changed = True
    if 'strikethrough' in changes:
        font.strike = bool(changes['strikethrough'])
        changed = True
    if 'font_family' in changes and _apply_font_family(font, changes['font_family']):
        changed = True
    if 'font_size' in changes:
        font.size = (
            _OPENPYXL_DEFAULT_FONT_SIZE
            if changes['font_size'] is None
            else changes['font_size']
        )
        changed = True
    if 'text_color' in changes:
        if changes['text_color'] is None:
            font.color = _OPENPYXL_DEFAULT_FONT_COLOR
        else:
            rgb = _excel_rgb(changes['text_color'])
            if not rgb:
                raise XlsxPatchExportError(
                    f'Unsupported text color {changes["text_color"]!r}.'
                )
            font.color = rgb
        changed = True
    if changed:
        excel_cell.font = font


def _patch_number_format(excel_cell, changes):
    if 'number_format' not in changes:
        return
    value = changes['number_format']
    if value is None:
        excel_cell.number_format = 'General'
        return
    mapped = _excel_number_format(value)
    excel_cell.number_format = mapped or 'General'


def _patch_format(worksheet, key, changes):
    excel_cell = _excel_cell(worksheet, key)
    _patch_font(excel_cell, changes)
    _patch_number_format(excel_cell, changes)


def _patch_highlight(worksheet, key, color):
    excel_cell = _excel_cell(worksheet, key)
    if color is None:
        excel_cell.fill = PatternFill(fill_type=None)
        return
    if not _excel_rgb(color):
        raise XlsxPatchExportError(f'Unsupported highlight color {color!r}.')
    _apply_fill(excel_cell, color)